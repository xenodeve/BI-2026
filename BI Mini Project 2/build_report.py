"""Build matching Thai PDF/Word scaffolds from raw audit and actual native evidence."""
from pathlib import Path
import argparse, csv, html, json, sys
from collect_evidence import ROOT, MODELS, METRICS, DROP, BATCH_ENGINE, analyze_raw, collect_native, create_templates, load_json

def metric_text(row, key):
    try:
        value = float(row[key])
        if not 0 <= value <= 1: raise ValueError()
        return f'{value:.4f}'
    except (KeyError, ValueError, TypeError): return 'PENDING'

def native_image_roi(path):
    """Display original native pixels using report cropping; never rewrite sources."""
    name = path.name
    if name.startswith('native_common_') and name.endswith('_performance.png'):
        return (0, 90, 640, 1008)
    return {
        'native_common_RF_importance.png': (108, 185, 430, 690),
        'native_common_DL_scoring.png': (480, 185, 1520, 265),
        'native_common_DL_confusion.png': (270, 190, 1510, 350),
        'native_common_DT_reading.png': (0, 90, 1020, 420),
        'native_common_LR_reading.png': (108, 185, 765, 415),
    }.get(name)

def sections(quality, status, rows, setup, reading):
    setup = setup or {}; reading = reading or {}
    blocks = []
    def heading(text): blocks.append(('heading', text))
    def para(text): blocks.append(('paragraph', text))
    def table(headers, data, widths): blocks.append(('table', (headers, data, widths)))
    def page(): blocks.append(('pagebreak', None))
    complete = status['status'] == 'COMPLETE_NATIVE_EVIDENCE'
    default_checkpoint = status.get('default_checkpoint_displayed', False)
    numeric_verified = status.get('common_holdout_numeric_verified', False)
    batch = setup.get('engine') == BATCH_ENGINE
    blocks.append(('title', 'การจำแนกผู้ป่วยนอกและผู้ป่วยที่เข้ารักษาในโรงพยาบาล'))
    para('Mini Project 2  Data Model and Evaluation ด้วย Altair AI Studio Auto Model')
    para('รายงานนี้ใช้ข้อมูล Covid Data.csv และ PATIENT_TYPE เป็น Target เพื่อศึกษาการจำแนกสองกลุ่มตามบทเรียน โดย Hospitalized เป็น Positive Class มีการวิเคราะห์ข้อมูลดิบจริงและพื้นที่สำหรับผล Auto Model ทั้ง 9 เทคนิค')
    para('สถานะ ' + ('มีหลักฐาน native ครบตามการตรวจไฟล์ ต้องอ่านควบคู่กับภาพและยืนยันว่าภาพเป็นผลจากการรันจริง' if complete else 'DRAFT  ผล native common comparison และการเลือกโมเดลตรวจได้แล้ว แต่ยังรอหลักฐาน GUI ที่ตรงกับ common run' if numeric_verified else 'PENDING NATIVE EVIDENCE  ยังไม่ใช่รายงานผลทดลองที่สมบูรณ์ ช่อง PENDING ไม่มีค่าทดแทนจาก Python'))
    if numeric_verified:
        para('ผลเปรียบเทียบปัจจุบันมาจาก official AI Studio batch ที่รัน native algorithms จริง ทุกโมเดลใช้ Training18,000 และ Testing12,000 ระเบียนเดียวกัน ตรวจ ID/Label/Prediction และ metrics ผ่านแล้ว เลือก Deep Learning จากผลเปรียบเทียบนี้ และบันทึก native model readings ที่ตรวจ fit เดียวกัน ภาพ GUI default เดิมใช้เป็นบริบทและใช้ Test ต่างชุด')
        para('ภาพ native GUI common PerformanceVector ทั้ง9โมเดลยืนยันตรงกับ batch comparison แล้ว' if status.get('common_gui_result_photos_verified') else 'ยังรอภาพ native GUI common PerformanceVector ที่ตรงกับผล batch')
    if default_checkpoint:
        para('Auto Model native รันสำเร็จและบันทึกครบ 9 โมเดลแล้ว ตัวเลขในรายงานเป็นผล checkpoint จริง แต่ generated process ใช้ sampling cap ต่างกันก่อนแบ่ง Train/Test จึงไม่ผ่านข้อกำหนด Testing Data ชุดเดียวกัน ขณะนี้รอ native common-holdout re-evaluation ผลเดิมแยกเก็บใน Native_Default_Model_Comparison.csv ไม่ใช่ตาราง final')
    heading('1 Dataset')
    para('ข้อมูลผู้รับบริการเกี่ยวกับ COVID-19 จาก Dataset ที่ระบุในแผนและบทเรียน URL ' + quality['source_url'] + ' ใช้ไฟล์จริงที่เก็บใน data/Covid Data.csv การตรวจนี้ไม่ดาวน์โหลดข้อมูลใหม่หรือรับรองเนื้อหาเว็บไซต์ปัจจุบัน')
    para(f"ต้นทาง {quality['raw_rows']:,} ระเบียน {quality['raw_columns']} คอลัมน์ การตัดเฉพาะคอลัมน์ 5 ตัวตามแผนคง {quality['planned_prepared_rows_without_dropping_records']:,} ระเบียน 16 คอลัมน์ รวม Target และ Input 15 ตัว จึงผ่านเกณฑ์หลังเตรียมอย่างน้อย 1,000,000 ระเบียนโดยไม่เพิ่มแถวเทียม จำนวน native ที่ใช้จริงต้องอ่านจาก Model Setup")
    table(['Target', 'ความหมาย', 'จำนวน', 'สัดส่วน'], [
        ['1', 'Outpatient', f"{quality['target_counts']['1']:,}", f"{1-quality['hospitalized_fraction']:.2%}"],
        ['2', 'Hospitalized positive', f"{quality['target_counts']['2']:,}", f"{quality['hospitalized_fraction']:.2%}"]], [45, 220, 110, 140])
    para('Target ต้องเปลี่ยนรหัส 1 เป็น Outpatient และ 2 เป็น Hospitalized ใน native preparation แล้วกำหนด Positive Class = Hospitalized ให้เหมือนกันทุกโมเดล ห้ามตีความรหัส 1/2 เป็นตัวแปรต่อเนื่อง')
    para('Input ที่ตกลงใช้ ' + ', '.join(quality['inputs']) + ' โดย AGE เป็น numeric ตัวแปรที่เหลือเป็น nominal ไม่มี ID ระบุผู้ป่วยในชุดนี้')
    para('ตัด PREGNANT ตามขั้นตอนบทเรียน และตัด DATE_DIED, INTUBED, ICU, CLASIFFICATION_FINAL ซึ่งไม่เหมาะเป็นข้อมูลทำนายก่อนทราบสถานะการเข้ารักษา ชื่อ CLASIFFICATION_FINAL เป็นการสะกดจริงใน CSV')
    para('ขอบเขตการตีความกำหนดว่าข้อมูล Input ต้องทราบก่อนตัดสินใจรับเข้ารักษา แต่ Dataset นี้ไม่มีเวลาเก็บแต่ละตัวแปรให้พิสูจน์ข้อสมมติฐานดังกล่าว โดยเฉพาะ PNEUMONIA, MEDICAL_UNIT และ USMER จึงเป็นงานเปรียบเทียบเพื่อการเรียน ไม่ใช่ระบบคัดกรองคลินิกที่ผ่าน validation หรือข้อสรุปเชิงเหตุและผล')
    para(f"ไม่พบช่องว่างจริง แต่มีรหัส 97/98/99 ในฟิลด์ nominal ต้องคงเป็นกลุ่มไม่ทราบหรือจัดการตามความหมายจริง ไม่แทนด้วยค่าเฉลี่ยของรหัส AGE อยู่ระหว่าง {quality['age']['min']} ถึง {quality['age']['max']} ปี ค่า AGE 97/98/99 เป็นอายุ ไม่ใช่รหัส missing โดยอัตโนมัติ")
    para(f"พบแถว raw ซ้ำทั้ง profile {quality['exact_duplicate_raw_rows']:,} แถว หลังตัดคอลัมน์ profile รวม Target ซ้ำ {quality['duplicate_prepared_profiles_including_target']:,} แถว และ Input ไม่รวม Target ซ้ำ {quality['duplicate_input_profiles_excluding_target']:,} แถว ตัวเลขนี้นับแถวหลังแถวแรกในแต่ละ profile ไม่พิสูจน์ว่าเป็นคนเดิม จึงไม่ลบอัตโนมัติและไม่ทำซ้ำข้อมูลเพื่อให้ครบเกณฑ์")
    page(); heading('2 Model Setup')
    para('ลำดับ native ที่ต้องมีหลักฐานคือ Turbo Prep ตรวจข้อมูลและตัดคอลัมน์ กำหนดชนิด nominal/numeric แปลง Target จากนั้น Auto Model > Predict เลือก Target และ Positive Class เลือก Inputs และทดสอบทั้ง 9 โมเดลด้วย Testing Data ชุดเดียวกัน')
    fields = [('prepared_rows', 'ระเบียนหลังเตรียม native'), ('model_rows_used', 'ระเบียนที่ Auto Model ใช้จริง'),
              ('training_rows', 'Training rows'), ('test_rows', 'Testing rows'), ('split_method', 'วิธีแบ่งข้อมูล'),
              ('test_set_id', 'TestSetID เดียวกัน'), ('seed', 'Seed ที่ native แสดง'),
              ('sampling_or_license_limit', 'ข้อจำกัดหรือการเลือกตัวอย่าง'), ('missing_code_policy', 'การจัดการรหัส missing'),
              ('preprocessing_fit_scope', 'ขอบเขต fit preprocessing')]
    table(['รายการ', 'ค่าจาก native'], [[label, str(setup.get(key)) if setup.get(key) is not None else 'PENDING'] for key, label in fields], [190, 325])
    para('หาก Auto Model เลือกตัวอย่างภายในหรือ license จำกัดข้อมูล ต้องระบุจำนวนที่ใช้จริง แยกจากจำนวนต้นทางและจำนวนหลัง Turbo Prep ไม่อนุมาน train/test split หรือ seed จากค่าเริ่มต้นของโปรแกรม')
    para('การแบ่งแบบสุ่มอาจมี profile เดียวกันอยู่ทั้ง Train และ Test และไม่มี patient ID สำหรับแยกคนละบุคคล ดังนั้นผลไม่ได้พิสูจน์ความสามารถกับผู้ป่วยใหม่หรือช่วงเวลาใหม่ การ fit preprocessing ต้องใช้ Training เท่านั้นและบันทึกขั้นตอนที่ Auto Model สร้างจริง')
    if numeric_verified:
        para('Common native run ใช้ sample30,000 เพื่อควบคุมต้นทุนคำนวณ หลัง Turbo Prep เก็บข้อมูลจริง1,048,575ระเบียนครบ Initial Sample local seed2001; Split Data60/40 local seed1992; ตรวจทุก12,000 Test rows โดยไม่ใช้ robust five-of-seven averaging ตัว audit COMMON_ROW_ID เป็น special ID ไม่ใช่ predictor')
    if default_checkpoint:
        audit=load_json('native_result_audit.json')
        para('Generated process native ระบุ root seed 2001 และ Split Data 60/40 local seed 1992 แต่ Sample ก่อน split ใช้ cap ต่างกัน จึงไม่ทำให้ Test เหมือนกัน Performance มาจาก 7 holdout batches แล้ว robust average คง 5 ค่า ไม่ใช่ 5-fold cross-validation จำนวน confusion รวมด้านล่างคือ 5 batches ที่คงไว้ ไม่ใช่จำนวนข้อมูลทั้งไฟล์หรือ full validation')
        table(['Model', 'Initial sampling cap', 'Counter N'], [[m,audit['models'][m]['process']['initial_sample_size_cap'],audit['models'][m]['confusion_matrix']['n']] for m in MODELS], [235,140,140])
    para('ภาพที่ต้องอ่านประกอบ Predict, Target/Positive Class, Inputs/ชนิดข้อมูล, Training/Testing และจำนวนแถว ภาคผนวกแนบเฉพาะไฟล์ PNG/JPEG ที่มีอยู่จริงใน screenshots ภาพ import อย่างเดียวไม่ยืนยันว่า Run สำเร็จ')
    page(); heading('3 Model Comparison')
    para('ตารางใช้ Accuracy, Precision, Recall, F1 และ AUC เป็นสัดส่วน 0 ถึง 1 คัดจาก Testing Data ใน native เท่านั้น AUC ต้องมาจาก confidence หรือ score ไม่ใช่ predicted class อย่างเดียว ชื่อ Engine และ Positive Class ต้องตรงกันทุกแถว')
    if default_checkpoint:
        para('ตารางนี้แสดง native DEFAULT checkpoint ที่ใช้ Test ต่างชุด เปรียบเทียบเชิงบันทึกผลได้ แต่ยังใช้เลือกผู้ชนะที่ยุติธรรมหรืออ้างว่าผ่าน shared-test requirement ไม่ได้ ค่าเฉลี่ยใช้ meanSum/averageCount และ sample SD บันทึกแยกใน CSV โดย crosscheck Accuracy = 1 - Classification Error และ SD กับ Native_Comparison_Raw.csv ผ่านครบ 9')
    by_model = {x.get('Model'): x for x in rows}
    table(['Model'] + METRICS, [[m] + [metric_text(by_model.get(m, {}), k) for k in METRICS] for m in MODELS], [140, 75, 75, 75, 75, 75])
    para(f"Baseline จาก raw คือทำนาย Outpatient ทุกแถว Accuracy {quality['majority_class_accuracy_baseline']:.4f} และ Recall ของ Hospitalized เท่ากับ 0 ดังนั้น Accuracy สูงเพียงอย่างเดียวไม่เพียงพอ ต้องอ่าน Recall/F1 และจำนวน False Negative ควบคู่ ส่วน baseline ของ Test ต้องคำนวณจากสัดส่วน Test จริงเมื่อมีข้อมูล")
    if complete or default_checkpoint or numeric_verified:
        maxima=[]
        for metric in METRICS:
            maximum = max(float(x[metric]) for x in rows)
            winners = ', '.join(x['Model'] for x in rows if abs(float(x[metric])-maximum) < 1e-10)
            maxima.append(f'{metric}: {winners} {maximum:.4f}')
        para('ค่าสูงสุดจาก native comparison  '+'; '.join(maxima)+(' แต่ Default checkpoint ใช้ Test ต่างกัน จึงยังไม่ใช่ผู้ชนะในการเปรียบเทียบตามโจทย์' if default_checkpoint else ''))
    else:
        para('ผู้ชนะของแต่ละ metric ยัง PENDING จนมีผลทั้ง 9 โมเดลและบริบท Test/Positive Class/ภาพครบ ไม่ใช้ผลของตัวอย่างอาจารย์หรือ Python proxy เติมตาราง')
    for model in MODELS:
        row = by_model.get(model)
        if row:
            para(model + '  ' + (row.get('Notes') or 'PENDING คำอธิบาย 2 ถึง 3 บรรทัดจาก metrics และข้อผิดพลาดของ native result'))
        else:
            para(model + '  PENDING ผล native และสรุปจุดเด่นหรือข้อด้อย 2 ถึง 3 บรรทัด')
    page(); heading('4 Model Reading')
    lr = reading.get('logistic_regression', {})
    if batch:
        para('อ่านเฉพาะ native artifacts ที่ยืนยันเป็น fit เดียวกับ common evaluation ผล coefficients/tree/importance จาก Default checkpoint เดิมไม่ถูกนำมาแทน common fit ส่วน Confusion Matrix ตรวจจาก common native predictions และภาพ PerformanceVector ที่ตรงกัน Scoring เป็นการเปิดระเบียนที่ native เคยทำนายแล้ว ไม่ใช่การเปลี่ยน threshold ใน Simulator')
        if status.get('same_evaluation_fit_model_artifacts_verified'):
            para('Instrumented full-sequence native replay บันทึก evaluated LR/DT/RF จาก model output ที่ใช้ทดสอบจริง ทั้ง36 replay exports รวม prediction/confidence ตรงกับ canonical แบบ byte-for-byte และ native process จบสำเร็จ จึงเชื่อม model readings กับ common fit นี้ได้ ไม่ใช้ Production Model ที่ฝึก30,000ระเบียนแทน')
    if default_checkpoint:
        para('การอ่านโมเดลในส่วนนี้มาจากไฟล์ native DEFAULT checkpoint เดิม เมื่อรัน common-holdout ใหม่ต้องอ่าน coefficients/tree/importance ของโมเดลใหม่อีกครั้งและเก็บภาพ GUI ให้ตรงกับไฟล์นั้น')
    para('Logistic Regression อ่าน weight โดยระบุ encoding และกลุ่มอ้างอิง ค่าบวกเพิ่ม score หรือ log odds ไปยัง Hospitalized ภายใต้ encoding ที่ native ใช้ ค่าลบไปทิศตรงข้าม น้ำหนักไม่ใช่เหตุและผล และเทียบขนาดตรง ๆ ไม่ได้เมื่อสเกลต่างกัน')
    if lr.get('note'):para(lr['note'])
    coef = lr.get('coefficients', [])
    table(['Feature', 'กลุ่มอ้างอิง', 'Weight', 'อ่านผล'], [[x.get('feature'), x.get('reference_category', ''), f"{x['weight']:.6f}" if isinstance(x.get('weight'),(int,float)) else 'PENDING', x.get('interpretation', '')] for x in coef] if coef else [['PENDING', '', 'PENDING', 'รอ coefficient จาก native Logistic Regression']], [130, 100, 70, 215])
    tree = reading.get('decision_tree', {})
    para('Decision Tree root ' + str(tree.get('root') or 'PENDING') + ' เส้นทางถึง leaf ' + (' > '.join(map(str, tree.get('path_to_leaf', []))) or 'PENDING') + ' ผล leaf ' + str(tree.get('leaf_prediction') or 'PENDING') + ' ต้องอ้างเงื่อนไข split และ leaf จากต้นไม้ native ไม่สร้างกฎจากการคาดเดา')
    if tree.get('leaf_training_counts'):para('จำนวนใน training leaf '+json.dumps(tree['leaf_training_counts'],ensure_ascii=False)+' เป็นจำนวนใน node ฝึก ไม่ใช่ Confusion Matrix ของ Test')
    forest = reading.get('random_forest', {}).get('top3_importance', [])
    table(['Rank', 'Feature', 'Importance'], [[i+1, x.get('feature'), x.get('importance')] for i, x in enumerate(forest)] if forest else [['PENDING', 'PENDING', 'PENDING']], [70, 245, 200])
    para('Importance อธิบายการใช้ตัวแปรภายในโมเดล ไม่ยืนยันสาเหตุหรือทิศทางผล ต้องระบุค่าตาม native และกลุ่มตัวแปรหลัง encoding')
    if reading.get('random_forest',{}).get('note'): para(reading['random_forest']['note'])
    cm = reading.get('confusion_matrix', {})
    table(['Confusion positive Hospitalized', 'Actual Hospitalized', 'Actual Outpatient'],
          [['Predicted Hospitalized', 'TP ' + str(cm.get('tp') if cm.get('tp') is not None else 'PENDING'), 'FP ' + str(cm.get('fp') if cm.get('fp') is not None else 'PENDING')],
           ['Predicted Outpatient', 'FN ' + str(cm.get('fn') if cm.get('fn') is not None else 'PENDING'), 'TN ' + str(cm.get('tn') if cm.get('tn') is not None else 'PENDING')]], [235, 140, 140])
    para('Recall = TP/(TP+FN) เป็น metric สำคัญเมื่อให้ความสำคัญกับการลด False Negative ต้องพิจารณา Precision/FP และต้นทุนด้วย การปรับ threshold ต้องใช้ชุด validation แยกและตรวจใหม่ ไม่เลือก threshold จาก Test ที่รายงาน')
    score = reading.get('scoring', {})
    para('Scoring ตัวอย่าง 1 ระเบียน Model ' + str(score.get('model') or 'PENDING') + ' Prediction ' + str(score.get('prediction') or 'PENDING') + ' Confidence Hospitalized ' + str(score.get('confidence_hospitalized') if score.get('confidence_hospitalized') is not None else 'PENDING') + ' Threshold 0.50 โดย prediction ที่ threshold นี้ต้องเทียบ confidence ของ Positive Class ไม่สลับกับ confidence ของ Outpatient')
    if score.get('record_inputs'):
        para('Inputs ของระเบียนจริง ' + json.dumps(score['record_inputs'], ensure_ascii=False))
    if score.get('threshold_interpretation'):para(score['threshold_interpretation'])
    page(); heading('5 Conclusion')
    if complete or numeric_verified:
        para('Final Model ที่บันทึก ' + setup['selected_model'] + ' เหตุผลจาก native evaluation ' + setup['selection_reason'])
        if not complete:para('การเลือกโมเดลและ common numeric comparison ผ่านแล้ว แต่ไฟล์นี้ยังเป็น Draft เพราะหลักฐานการอ่านโมเดลหรือภาพ common-run GUI ยังไม่ครบ ไม่อ้างว่าการเลือกจาก Test เดียวกันเป็น independent validation หลังเลือก')
    else:
        para('Final Model ยัง PENDING ต้องมีการเปรียบเทียบ Test เดียวกันและหลักฐานการอ่านโมเดลพอให้เลือก โดยพิจารณา metric ทั้ง 5 ข้อผิดพลาดของ Test และเป้าหมายที่ให้ความสำคัญกับ Hospitalized')
        if default_checkpoint:
            para('ผล metric native ครบ 9 โมเดลแล้ว สิ่งที่ยังไม่ผ่านคือการใช้ Test เดียวกันและหลักฐาน final model/scoring ไม่ใช่การขาดผลจาก Auto Model เดิม ควรใช้ common-holdout native ใหม่โดยยืนยัน row ID และ label hash ตรงกันทุกโมเดลก่อนเลือกผลสุดท้าย')
    para('ข้อจำกัดคือข้อมูล anonymous profile ซ้ำจำนวนมาก ไม่มีเวลาของ input และไม่มีรหัสบุคคลสำหรับแบ่งผู้ป่วยหรือเวลาอย่างเป็นอิสระ Class imbalance ทำให้ Accuracy มี baseline สูง ผลที่รายงานต้องอ้าง split และ native preprocessing จริง และไม่อ้างความเหมาะสมใช้ตัดสินใจทางคลินิกจากงานเรียนนี้')
    heading('หลักฐานที่ยังขาดหรือไม่ผ่าน')
    if status['pending'] or status['errors']:
        for item in status['errors']:para('FAILED '+item)
        if batch and status['pending'] and all('screenshot' in x.lower() for x in status['pending']):
            for item in status['pending']:para('PENDING '+item)
            para('ตัวเลขและ fitted artifact ผ่านแล้ว ภาพ Default GUI เดิมใช้บริบทอย่างเดียว ไม่ใช้แทนภาพ common-result การตรวจ final report ยังไม่ผ่านจนมีภาพที่ตรงกับ test identity และ artifact จริง')
        else:
            for item in status['pending']:para('PENDING '+item)
    else:
        para('การตรวจโครงสร้างหลักฐานผ่าน ไม่มีช่อง pending ใน collector ยังต้องตรวจความถูกต้องของการถอดค่ากับ GUI และภาพจริงโดยผู้ควบคุมงาน')
    para('แหล่งข้อกำหนด BI Mini Project 2.pdf หน้า 1 ถึง 3 และแผน Task Mini Project 2 บทเรียน BI08 เป็น PDF ภาพสแกน text extract ว่าง จึงไม่ใช้ข้อความ extract ว่างอ้างว่าตรวจทุกหน้าผ่าน')
    para('ตรวจสอบซ้ำด้วย collect_evidence.py และสร้างรายงานใหม่ด้วย build_report.py เมื่อมี native_model_comparison.csv, native_setup.json, native_model_reading.json และ screenshots จริง')
    page(); heading('ภาคผนวกคุณภาพข้อมูล')
    para('จำนวนรหัสที่ตรวจจาก raw จริง เป็นการตรวจข้อมูลด้วยไฟล์ ไม่ใช่หลักฐานว่าได้กด Turbo Prep รหัสในคอลัมน์ AGE ไม่ถูกรวมเป็น missing')
    data = [[c, x['97'], x['98'], x['99'], 'ตัดออก' if c in DROP else 'คงไว้'] for c, x in quality['missing_code_counts'].items() if sum(x.values())]
    table(['Attribute', '97', '98', '99', 'Role'], data, [175, 85, 85, 85, 85])
    para('SHA256 ' + quality['source_sha256'])
    screenshots = sorted(p for p in (ROOT/'screenshots').iterdir() if p.suffix.lower() in ('.png', '.jpg', '.jpeg') and not any(x in p.name.lower() for x in ['_error','_probe']))
    for screenshot in screenshots:
        page(); heading('ภาคผนวกภาพหลักฐานที่มีจริง')
        if batch and screenshot.name in ['native_05_nine_models.png','native_06_all9_completed.png','native_NB_performance.png']:
            context='Context Unequal Test  ภาพ original GUI Auto Model default เดิม ใช้ Test ต่างชุด ไม่ใช่ภาพผล official batch common-holdout ในตารางปัจจุบัน'
        elif batch and screenshot.name.startswith(('native_01','native_02','native_03','native_04')):
            context='ภาพ original GUI data preparation และ target/input setup ข้อมูล1,048,575ระเบียน ไม่ใช่ภาพผล common-holdout batch'
        elif screenshot.name.startswith('native_common_') and screenshot.name.endswith('_performance.png'):
            context='ภาพ native GUI full common PerformanceVector ทั้ง12,000ระเบียน ตรงกับผล official batch ในตาราง; ภาพในรายงานตัดขอบเพื่อให้อ่านตัวเลขได้ ไฟล์ต้นฉบับเต็มไม่แก้ไข'
        elif screenshot.name=='native_common_RF_importance.png':
            context='ภาพ native viewer ของ genuine RF learner weights ที่บันทึกจาก evaluated common fit; ตัดขอบแสดงตารางต้นฉบับ ไม่สร้าง UI ใหม่'
        elif screenshot.name=='native_common_DL_scoring.png':
            context='ภาพ native viewer ของ saved scored row348; threshold0.50เป็นการคำนวณในรายงาน ไม่ใช่ live Simulator; ตัดขอบเฉพาะค่าจากภาพต้นฉบับ'
        elif screenshot.name=='native_common_DL_confusion.png':
            context='ภาพ native full common DL PerformanceVector; rows predicted/columns true; ตัดขอบเฉพาะตารางต้นฉบับ'
        elif screenshot.name=='native_common_DT_reading.png':
            context='ภาพ native Tree Description ของ saved evaluated common fit; root PNEUMONIA=1 และ leaf Hospitalized343/1735เป็น training-node counts; ภาพตัดขอบเฉพาะส่วนต้นของต้นไม้'
        elif screenshot.name=='native_common_LR_reading.png':
            context='ภาพ native ReadCSV ของ coefficient export ที่ตรงกับ arrays ใน saved evaluated LR model ทุกค่า ไม่อ้างว่าเป็น coefficient tab ภายใน model UI; ภาพตัดขอบเฉพาะตารางต้นฉบับ'
        elif screenshot.name=='native_saved_LR_description_context.png':
            context='Context only: native LR Description/training history ไม่แสดง coefficients และไม่ใช้แทนหลักฐาน coefficient table'
        else:context='ภาพนี้แสดงเฉพาะสิ่งที่ปรากฏในภาพ ไม่ทำให้ข้อกำหนดอื่นผ่านโดยอัตโนมัติ'
        para(screenshot.name + '  '+context)
        blocks.append(('image', screenshot))
    return blocks

def make_pdf(blocks, path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image, Flowable
    from PIL import Image as PILImage
    class NativeClip(Flowable):
        def __init__(self, source, roi):
            Flowable.__init__(self)
            self.source, self.roi = source, roi
            self.original_width, self.original_height = PILImage.open(source).size
            left, top, right, bottom = roi
            self.scale = min(515/(right-left), 620/(bottom-top))
            self.width, self.height = (right-left)*self.scale, (bottom-top)*self.scale
        def draw(self):
            c = self.canv; c.saveState()
            clip = c.beginPath(); clip.rect(0, 0, self.width, self.height); c.clipPath(clip, stroke=0)
            left, top, right, bottom = self.roi
            c.drawImage(str(self.source), -left*self.scale, (bottom-self.original_height)*self.scale,
                        width=self.original_width*self.scale, height=self.original_height*self.scale)
            c.restoreState()
    pdfmetrics.registerFont(TTFont('Thai', r'C:\Windows\Fonts\tahoma.ttf', shapable=True))
    pdfmetrics.registerFont(TTFont('ThaiBold', r'C:\Windows\Fonts\tahomabd.ttf', shapable=True))
    body = ParagraphStyle('Body', fontName='Thai', fontSize=10, leading=16, spaceAfter=7, wordWrap='CJK', shaping=1)
    heading = ParagraphStyle('Heading', parent=body, fontName='ThaiBold', fontSize=14, leading=21, spaceAfter=10, keepWithNext=True)
    title = ParagraphStyle('Title', parent=heading, fontSize=20, leading=29, spaceAfter=15)
    cell = ParagraphStyle('Cell', parent=body, fontSize=8.4, leading=13, spaceAfter=0)
    head = ParagraphStyle('Head', parent=cell, fontName='ThaiBold', textColor=colors.white)
    p = lambda text, style=body: Paragraph(html.escape(str(text)), style)
    story = []
    for kind, value in blocks:
        if kind in ('paragraph', 'heading', 'title'): story.append(p(value, {'paragraph': body, 'heading': heading, 'title': title}[kind]))
        elif kind == 'pagebreak': story.append(PageBreak())
        elif kind == 'table':
            headers, rows, widths = value
            table = Table([[p(x, head) for x in headers]] + [[p(x, cell) for x in row] for row in rows], colWidths=widths, repeatRows=1)
            commands = [('BACKGROUND', (0,0), (-1,0), colors.HexColor('#203A5F')), ('GRID', (0,0), (-1,-1), .4, colors.HexColor('#D9D9D9')), ('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('LEFTPADDING', (0,0), (-1,-1), 6), ('RIGHTPADDING', (0,0), (-1,-1), 6), ('TOPPADDING', (0,0), (-1,-1), 6), ('BOTTOMPADDING', (0,0), (-1,-1), 6)]
            for i in range(2, len(rows)+1, 2): commands.append(('BACKGROUND', (0,i), (-1,i), colors.HexColor('#F0F4F8')))
            table.setStyle(TableStyle(commands)); story.extend([table, Spacer(1, 10)])
        elif kind == 'image':
            roi = native_image_roi(value)
            if roi:story.append(NativeClip(value, roi))
            else:
                w, h = PILImage.open(value).size; scale = min(515/w, 675/h)
                story.append(Image(str(value), width=w*scale, height=h*scale))
    def footer(canvas, doc):
        canvas.setFont('Thai', 8); canvas.drawString(40, 22, 'Mini Project 2   ' + ('PENDING NATIVE EVIDENCE' if 'Draft' in path.name else 'Native evidence report')); canvas.drawRightString(555, 22, str(doc.page))
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=40, rightMargin=40, topMargin=36, bottomMargin=38).build(story, onFirstPage=footer, onLaterPages=footer)

def make_docx(blocks, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(); sec = doc.sections[0]; sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.6)
    for name in ['Normal', 'Title', 'Heading 1']:
        style=doc.styles[name];style.font.name='Tahoma';style.font.color.rgb=RGBColor(0,0,0)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Tahoma')
    doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.line_spacing=1.25
    doc.styles['Normal'].paragraph_format.space_after=Pt(7);doc.styles['Title'].font.size=Pt(20)
    doc.styles['Heading 1'].font.size=Pt(14)
    for kind,value in blocks:
        if kind in ('paragraph','heading','title'):doc.add_paragraph(str(value),style={'paragraph':'Normal','heading':'Heading 1','title':'Title'}[kind])
        elif kind=='pagebreak':doc.add_page_break()
        elif kind=='image':
            from PIL import Image
            w,h=Image.open(value).size;roi=native_image_roi(value)
            if roi:
                left,top,right,bottom=roi;scale=min(6.9/(right-left),8.3/(bottom-top))
                picture=doc.add_picture(str(value),width=Inches((right-left)*scale),height=Inches((bottom-top)*scale))
                crop=OxmlElement('a:srcRect')
                for edge,amount in [('l',left/w),('t',top/h),('r',(w-right)/w),('b',(h-bottom)/h)]:crop.set(edge,str(round(amount*100000)))
                picture._inline.graphic.graphicData.pic.blipFill.insert(1,crop)
            else:
                scale=min(6.9/w,9.0/h);doc.add_picture(str(value),width=Inches(w*scale),height=Inches(h*scale))
        elif kind=='table':
            headers,rows,widths=value;t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            for i,width in enumerate(widths):t.columns[i].width=Inches(width/515*6.9)
            for j,x in enumerate(headers):t.rows[0].cells[j].text=str(x)
            for row in rows:
                cells=t.add_row().cells
                for j,x in enumerate(row):cells[j].text=str(x)
            for i,row in enumerate(t.rows):
                for j,c in enumerate(row.cells):
                    c.width=Inches(widths[j]/515*6.9);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcpr.append(borders);margins=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
                    tcpr.append(margins)
                    fill='203A5F' if i==0 else 'F0F4F8' if i%2==0 else 'FFFFFF'
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),fill);tcpr.append(shade)
                    for p in c.paragraphs:
                        p.paragraph_format.space_after=Pt(2)
                        for run in p.runs:
                            run.font.size=Pt(8.4)
                            if i==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
    sec.footer.paragraphs[0].text='Mini Project 2   ' + ('PENDING NATIVE EVIDENCE' if 'Draft' in path.name else 'Native evidence report')
    doc.save(path)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--require-complete',action='store_true');args=parser.parse_args()
    create_templates()
    quality=load_json('raw_data_quality.json') or analyze_raw()
    status,rows,setup,reading=collect_native()
    if args.require_complete and (status['pending'] or status['errors']):
        print('REFUSED FINAL REPORT: native evidence is incomplete or inconsistent');return 2
    if not rows and (ROOT/'Native_Default_Model_Comparison.csv').exists():
        rows=list(csv.DictReader((ROOT/'Native_Default_Model_Comparison.csv').open(encoding='utf-8-sig')))
        setup=load_json('native_default_setup.json');reading=load_json('native_default_model_reading.json')
        status['default_checkpoint_displayed']=True
        status['pending'].append('Native DEFAULT checkpoint used unequal test sets; common-holdout results are pending and no final model is selected.')
    blocks=sections(quality,status,rows,setup,reading)
    stem='MiniProject2_Report' if status['status']=='COMPLETE_NATIVE_EVIDENCE' else 'MiniProject2_Report_Draft'
    make_pdf(blocks,ROOT/(stem+'.pdf'));make_docx(blocks,ROOT/(stem+'.docx'))
    print(stem+'.pdf and .docx created; Word pagination requires separate native rendering.')
    return 0

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8');raise SystemExit(main())
