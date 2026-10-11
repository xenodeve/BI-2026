from pathlib import Path
from zipfile import ZipFile
import json,hashlib,math,xml.etree.ElementTree as ET
from openpyxl import load_workbook
from pypdf import PdfReader
ROOT=Path(r'D:\Github\BI 2026');W2=ROOT/'Workshop 2 Data Presentation';W3=ROOT/'Workshop 3 What if Analysis in Excel'
v=json.loads((W3/'support/functional_completion.json').read_text(encoding='utf-8'));orig=W3/'Workshop3_WhatIf.xlsx';completed=W3/'Workshop3_WhatIf_Completed.xlsx';backup=Path(v['exact_current_backup'])
assert orig.read_bytes()==backup.read_bytes();assert hashlib.sha256(orig.read_bytes()).hexdigest()==v['original_sha256']
with ZipFile(completed)as z:
 for name in z.namelist():
  if name.endswith(('.xml','.rels')):ET.fromstring(z.read(name))
 assert len(z.namelist())==len(set(z.namelist()))
with ZipFile(orig)as before,ZipFile(completed)as after:
 for name in ['xl/styles.xml','xl/theme/theme1.xml','xl/sharedStrings.xml']:assert before.read(name)==after.read(name)
 for name in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml']:
  def formulas(z):
   tree=ET.fromstring(z.read(name));ns={'s':tree.tag.split('}')[0][1:]};return{c.attrib['r']:c.find('s:f',ns).text for c in tree.findall('.//s:c',ns)if c.find('s:f',ns)is not None}
  assert formulas(before)==formulas(after)
f=load_workbook(completed,data_only=False);d=load_workbook(completed,data_only=True)
assert f.sheetnames==['Goal Seek','Scenario','Scenario Summary'];sc=f['Scenario'].scenarios.scenario;assert len(sc)==4
assert math.isclose(d['Goal Seek']['B5'].value,425531.92056985485,abs_tol=1e-8);assert math.isclose(d['Goal Seek']['B16'].value,10000,abs_tol=1e-8)
for i,(s,c)in enumerate(zip(sc,v['scenario_cases'])):
 assert s.name==c['name'] and s.count==2
 assert [(x.r,float(x.val))for x in s.inputCells]==[('B5',float(c['price'])),('B8',float(c['years']))]
 expected=c['principal']/(sum((1+.07/12)**(-k)for k in range(1,c['years']*12+1)))
 assert abs(expected-c['payment'])<1e-8
 assert abs(d['Scenario Summary'][f'G{i+6}'].value-expected)<1e-8
 assert f['Scenario Summary'][f'G{i+6}'].value.startswith('=IF(')
files=[W2/'Workshop2_Data_Presentation.xlsx',completed,W2/'Workshop2_Data_Presentation_Report.pdf',W3/'Workshop3_WhatIf_Report.pdf',W2/'README.md',W3/'README.md'];outputs=[]
for file in files:
 item={'path':str(file),'bytes':file.stat().st_size,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()}
 if file.suffix=='.pdf':
  item['pages']=len(PdfReader(file).pages);assert len(list((file.parent/'qa/pdf').glob('page-*.png')))==item['pages']
 outputs.append(item)
for folder,count in [(W2,4),(W3,5)]:assert len(PdfReader(next(folder.glob('*Report.pdf'))).pages)==count
record={'result':'passed','all_9_PDF_pages_visually_reviewed':True,'native_GUI_claims':'only root-provided original 400k/9400 baseline; GoalSeek/ScenarioShow/Summary/CompletedReopen pending','original_and_fresh_backup_byte_identical':True,'genuine_scenarios':4,'original_formulas_unchanged':True,'original_styles_theme_sharedStrings_byte_identical':True,'outputs':outputs}
(W3/'support/final_delivery_verification.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(record,ensure_ascii=False,indent=2))

