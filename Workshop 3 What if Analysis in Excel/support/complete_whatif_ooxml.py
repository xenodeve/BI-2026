from pathlib import Path
import zipfile,xml.etree.ElementTree as ET,re,math,json,hashlib,shutil,datetime
from xml.sax.saxutils import escape
import openpyxl
from openpyxl.worksheet.scenario import InputCells,Scenario,ScenarioList

ROOT=Path(r'D:\Github\BI 2026');FOLDER=ROOT/'Workshop 3 What if Analysis in Excel';PATH=FOLDER/'Workshop3_WhatIf.xlsx';SUP=FOLDER/'support';SUP.mkdir(exist_ok=True)
STAMP=datetime.datetime.now().strftime('%Y%m%d_%H%M%S');BACKUP=FOLDER/'backups'/('Workshop3_WhatIf_current_before_functional_'+STAMP+'.xlsx');BACKUP.parent.mkdir(exist_ok=True);shutil.copy2(PATH,BACKUP)
original_hash=hashlib.sha256(PATH.read_bytes()).hexdigest();assert hashlib.sha256(BACKUP.read_bytes()).hexdigest()==original_hash
with zipfile.ZipFile(PATH)as z:original={n:z.read(n)for n in z.namelist()};infos={i.filename:i for i in z.infolist()}
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main';NS={'s':S}
def parsed_cell(xml,address):return ET.fromstring(xml).find(f'.//s:c[@r="{address}"]',NS)
def value(xml,address):return float(parsed_cell(xml,address).find('s:v',NS).text)
def style(xml,address):return parsed_cell(xml,address).get('s','0')
goal=original['xl/worksheets/sheet1.xml'].decode('utf-8');sc=original['xl/worksheets/sheet2.xml'].decode('utf-8');assert '<scenarios'not in sc
before={'B5':value(goal,'B5'),'B16':value(goal,'B16'),'observed_native_display_PMT':9400,'origin':'Root verified native PMT at starting price; caches preserved in exact backup'}
assert before['B5']==400000 and abs(before['B16']-9399.999874611907)<1e-8
down=value(goal,'B6');annual=value(goal,'B7');years=value(goal,'B8');target=value(goal,'B9');r=annual/12;periods=years*12
price=target*(-math.expm1(-periods*math.log1p(r)))/((1-down)*r) if r else target*periods/(1-down)
def payment(principal,annual,years):
    rate=annual/12;n=years*12
    return principal*rate/(-math.expm1(-n*math.log1p(rate))) if rate else principal/n
assert abs(price-425531.920569855)<1e-8
new_values={'B5':price,'B11':price*down,'B12':price*(1-down),'B16':payment(price*(1-down),annual,years)};new_values['B17']=new_values['B16']-target
def alter_value(xml,address,v):
    pattern=r'(<c\b[^>]*\br="'+address+r'"[^>]*>)(.*?)(</c>)'
    def f(m):
        content=m.group(2);assert '<v>'in content,(address,content);content=re.sub(r'<v>.*?</v>','<v>'+repr(v)+'</v>',content,count=1,flags=re.S);return m.group(1)+content+m.group(3)
    xml,count=re.subn(pattern,f,xml,count=1,flags=re.S);assert count==1,address;return xml
def alter_text(xml,address,text):
    pattern=r'(<c\b[^>]*\br="'+address+r'"[^>]*>)(.*?)(</c>)'
    def f(m):
        opening=re.sub(r'\s+t="[^"]*"','',m.group(1))[:-1]+' t="inlineStr">';return opening+'<is><t>'+escape(text)+'</t></is>'+m.group(3)
    xml,count=re.subn(pattern,f,xml,count=1,flags=re.S);assert count==1,address;return xml
for cell,v in new_values.items():goal=alter_value(goal,cell,v)
goal=alter_text(goal,'A30','ก่อนแก้: Native B5=400,000 และ PMT≈9,400. ตั้ง B5 เป็น425,531.920569855ด้วยสูตร PV ทางคณิตศาสตร์; ยังไม่ได้ใช้ native Goal Seek dialog')
goal=alter_text(goal,'A27','ผลB5ปัจจุบันเป็น mathematical solver. เมื่อเปิด Excel ให้ทดสอบ Goal Seek: Set B16, To10000, Change B5 แล้วเก็บหลักฐานจริง')
cases=[('200k_15y',200000,15),('200k_30y',200000,30),('300k_15y',300000,15),('300k_30y',300000,30)]
definitions=ScenarioList(current=0,show=0,sqref='B5 B8',scenario=[Scenario(name=name,count=2,locked=False,hidden=False,comment='Authored SpreadsheetML definition; native Show and Summary dialog not yet verified',inputCells=[InputCells(r='B5',val=str(p)),InputCells(r='B8',val=str(y))])for name,p,y in cases])
fragment=ET.tostring(definitions.to_tree(),encoding='unicode');assert fragment.startswith('<scenarios');sc=sc.replace('</sheetData>','</sheetData>'+fragment,1)
sc=alter_text(sc,'A37','มี4Scenario Manager definitionsในไฟล์แล้ว (SpreadsheetML B5,B8); ยังไม่ได้ native Show หรือสร้าง Summary dialog. ชีต Scenario Summary เป็น authored comparison')
sc=alter_text(sc,'D16','มี4Scenario Manager entriesแล้ว:200k_15y/200k_30y/300k_15y/300k_30y. Result B16; Native Show/Summaryยังรอตรวจ')
sc=alter_text(sc,'A34','ตารางsnapshotเดิมคงไว้. ชีตScenario Summaryใหม่มีสูตรที่linkเงินดาวน์/อัตราจากB6/B7และคำนวณ4กรณี; ไม่ใช่Summarydialog-generated')

# New worksheet reuses existing styles; it does not rewrite the original styles or themes.
money=style(original['xl/worksheets/sheet2.xml'],'B25');pct=style(original['xl/worksheets/sheet2.xml'],'B13');integer=style(original['xl/worksheets/sheet2.xml'],'C25');input_money=style(original['xl/worksheets/sheet2.xml'],'B5');input_year=style(original['xl/worksheets/sheet2.xml'],'B8');header=style(original['xl/worksheets/sheet2.xml'],'A24');title_style=style(original['xl/worksheets/sheet2.xml'],'A2');note_style=style(original['xl/worksheets/sheet2.xml'],'A34')
def cell(a,v=None,s='1',formula=None):
    if formula is not None:return f'<c r="{a}" s="{s}"><f>{escape(formula)}</f><v>{repr(v)}</v></c>'
    if isinstance(v,str):return f'<c r="{a}" s="{s}" t="inlineStr"><is><t>{escape(v)}</t></is></c>'
    return f'<c r="{a}" s="{s}"><v>{repr(v)}</v></c>'
rows=[]
rows.append('<row r="2" ht="28" customHeight="1">'+cell('A2','Scenario Summary (authored comparison)',title_style)+'</row>')
rows.append('<row r="3" ht="38" customHeight="1">'+cell('A3','ตารางเปรียบเทียบเขียนเองพร้อมสูตรและcached values; ไม่ใช่ native Scenario Summary dialog-generated',note_style)+'</row>')
headers=['สถานการณ์','ราคาบ้าน (บาท)','ปี','เงินดาวน์','ดอกเบี้ย/ปี','เงินกู้ (บาท)','ค่างวด/เดือน (บาท)','ดอกเบี้ยรวม (บาท)']
rows.append('<row r="5" ht="38" customHeight="1">'+''.join(cell(chr(65+i)+'5',h,header)for i,h in enumerate(headers))+'</row>')
result_cases=[];sd=value(sc,'B6');sa=value(sc,'B7');assert sd==.2 and sa==.07
for row,(name,p,y)in enumerate(cases,6):
    principal=p*(1-sd);pmt=payment(principal,sa,y);interest=pmt*y*12-principal
    result_cases.append({'name':name,'price':p,'years':y,'down':sd,'annual_rate':sa,'principal':principal,'payment':pmt,'total_interest':interest})
    f=f'IF(COUNT(B{row}:C{row},\'Scenario\'!$B$6:$B$7)<4,"ระบุ input ให้ครบ",IF(OR(B{row}<0,C{row}<=0,D{row}<0,D{row}>=1,E{row}<0),"input ไม่ถูกต้อง",B{row}*(1-D{row})))'
    values=[cell(f'A{row}',name),cell(f'B{row}',p,input_money),cell(f'C{row}',y,input_year),cell(f'D{row}',sd,pct,"'Scenario'!$B$6"),cell(f'E{row}',sa,pct,"'Scenario'!$B$7"),cell(f'F{row}',principal,money,f),cell(f'G{row}',pmt,money,f'IF(ISNUMBER(F{row}),-PMT(E{row}/12,C{row}*12,F{row},0,0),F{row})'),cell(f'H{row}',interest,money,f'IF(ISNUMBER(G{row}),G{row}*C{row}*12-F{row},G{row})')]
    rows.append(f'<row r="{row}" ht="26" customHeight="1">'+''.join(values)+'</row>')
for rn,text in [(12,'Inputsราคา/ปีในตารางเป็น4กรณีที่กำหนด. เงินดาวน์และดอกเบี้ยlinkจากScenario!B6/B7 จึงเปลี่ยนผลคำนวณทั้ง4แถวเมื่อแก้ assumptions.'),(14,'Scenario Manager definitionsใช้Scenario!B5,B8; native Show และSummarydialogยังรอตรวจในExcel. ไม่มีการใช้COM/VBA/GUIเพื่อสร้างหลักฐาน.'),(16,'ลดต้นลดดอก: rateรายปี/12, nper=ปี×12, จ่ายสิ้นเดือน(type0), fv0. PMTไม่รวมเงินดาวน์ ค่าธรรมเนียมหรือประกัน.')]:rows.append(f'<row r="{rn}" ht="40" customHeight="1">'+cell('A'+str(rn),text,note_style)+'</row>')
widths=[22,18,9,12,13,19,23,24]
sheet3='<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="'+S+'"><sheetPr><tabColor rgb="FF4778AE"/></sheetPr><dimension ref="A1:H17"/><sheetViews><sheetView showGridLines="0" workbookViewId="0"/></sheetViews><sheetFormatPr defaultRowHeight="24"/><cols>'+''.join(f'<col min="{i+1}" max="{i+1}" width="{w}" customWidth="1"/>'for i,w in enumerate(widths))+'</cols><sheetData>'+''.join(rows)+'</sheetData><mergeCells count="5">'+''.join(f'<mergeCell ref="A{n}:H{n}"/>'for n in [2,3,12,14,16])+'</mergeCells><pageMargins left="0.3" right="0.3" top="0.5" bottom="0.5" header="0.2" footer="0.2"/><pageSetup orientation="landscape" fitToWidth="1" fitToHeight="1"/></worksheet>'
parts=dict(original);parts['xl/worksheets/sheet1.xml']=goal.encode();parts['xl/worksheets/sheet2.xml']=sc.encode();parts['xl/worksheets/sheet3.xml']=sheet3.encode()
workbook=original['xl/workbook.xml'].decode();workbook=workbook.replace('</sheets>','<sheet name="Scenario Summary" sheetId="3" r:id="rId7"/></sheets>',1);workbook=re.sub(r'<calcPr\b[^>]*/>','<calcPr calcId="191029" calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>',workbook,1);parts['xl/workbook.xml']=workbook.encode()
rels=original['xl/_rels/workbook.xml.rels'].decode();rels=re.sub(r'<Relationship\b[^>]*Type="[^"]*/calcChain"[^>]*/>','',rels);rels=rels.replace('</Relationships>','<Relationship Id="rId7" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet3.xml"/></Relationships>');parts['xl/_rels/workbook.xml.rels']=rels.encode()
types=original['[Content_Types].xml'].decode();types=re.sub(r'<Override\b[^>]*PartName="/xl/calcChain.xml"[^>]*/>','',types);types=types.replace('</Types>','<Override PartName="/xl/worksheets/sheet3.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>');parts['[Content_Types].xml']=types.encode();parts.pop('xl/calcChain.xml',None)
app=original['docProps/app.xml'].decode();app=app.replace('<vt:i4>2</vt:i4>','<vt:i4>3</vt:i4>',1);app=app.replace('<TitlesOfParts><vt:vector size="2"','<TitlesOfParts><vt:vector size="3"',1);app=app.replace('<vt:lpstr>Scenario</vt:lpstr>','<vt:lpstr>Scenario</vt:lpstr><vt:lpstr>Scenario Summary</vt:lpstr>',1);parts['docProps/app.xml']=app.encode()
# Parse every XML part, including native namespace declarations kept byte-wise intact.
for name,data in parts.items():
    if name.endswith('.xml')or name.endswith('.rels'):ET.fromstring(data)
staging=FOLDER/'Workshop3_WhatIf_functional_staging.xlsx'
with zipfile.ZipFile(staging,'w',compression=zipfile.ZIP_DEFLATED)as z:
    for name,data in parts.items():z.writestr(infos[name] if name in infos else name,data)
check=openpyxl.load_workbook(staging,data_only=False);cached=openpyxl.load_workbook(staging,data_only=True)
assert check.sheetnames==['Goal Seek','Scenario','Scenario Summary'];s=check['Scenario'].scenarios;assert len(s.scenario)==4 and s.sqref=='B5 B8'
for scenario,(name,p,y)in zip(s.scenario,cases):
    assert scenario.name==name and scenario.count==2 and [(x.r,x.val)for x in scenario.inputCells]==[('B5',str(p)),('B8',str(y))]
for title in ['Goal Seek','Scenario']:
    old=openpyxl.load_workbook(BACKUP,data_only=False)[title]
    for row in old:
        for c in row:
            if c.data_type=='f':assert check[title][c.coordinate].value==c.value,(title,c.coordinate)
assert abs(cached['Goal Seek']['B16'].value-target)<1e-8
for rn,case in enumerate(result_cases,6):
    assert abs(cached['Scenario Summary'][f'G{rn}'].value-case['payment'])<1e-8
    assert 'PMT' in check['Scenario Summary'][f'G{rn}'].value
    # Independent geometric-sum formulation checks the annuity equation.
    q=1/(1+sa/12);annuity=sum(q**i for i in range(1,case['years']*12+1));assert abs(case['principal']/annuity-case['payment'])<1e-8
changed=[name for name in original if name not in parts or parts[name]!=original[name]];assert set(changed)=={'xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml','xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml','docProps/app.xml','xl/calcChain.xml'}
assert hashlib.sha256(PATH.read_bytes()).hexdigest()==original_hash,'Source changed during patch; refuse overwrite'
try:staging.replace(PATH);final=PATH
except PermissionError:
    final=FOLDER/'Workshop3_WhatIf_Completed.xlsx';staging.replace(final)
record={'result':'passed','final_workbook':str(final),'exact_current_backup':str(BACKUP),'original_sha256':original_hash,'before':before,'mathematical_goal_solution':{'price':price,'payment':new_values['B16'],'method':'PV/annuity mathematical solve; NOT native Goal Seek Run'},'native_GoalSeek_run':False,'native_ScenarioManager_Show_verified':False,'native_Summary_dialog_generated':False,'genuine_SpreadsheetML_scenarios':4,'changing_cells':['Scenario!B5','Scenario!B8'],'scenario_cases':result_cases,'authored_summary':'Scenario Summary: live PMT formulas + calculated caches, not native Summary dialog output','original_formulas_and_guards_unchanged':True,'original_styles_themes_sharedStrings_unchanged':True,'unrelated_part_bytes_preserved':True,'changed_original_parts':changed,'calcChain_removed_for_safe_rebuild':True,'checked_at':datetime.datetime.now().isoformat()}
(SUP/'functional_completion.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(record,ensure_ascii=False,indent=2))
